import os
import pandas as pd
from PyPDF2 import PdfReader
from docx import Document
from pptx import Presentation

def list_all_files_summary(root_folder):
    def get_file_extension(file):
        return os.path.splitext(file)[1].lower()  # Get the file extension

    def get_file_size(file_path):
        return os.path.getsize(file_path) / 1024  # Size in KB

    def get_pdf_page_count(file_path):
        try:
            with open(file_path, 'rb') as f:
                pdf = PdfReader(f)
                return len(pdf.pages)
        except Exception as e:
            print(f"Error reading {file_path}: {e}")
            return None

    def get_docx_page_count(file_path):
        try:
            doc = Document(file_path)
            paragraphs = len(doc.paragraphs)
            # Estimate page count based on paragraphs
            return max(1, paragraphs // 40)  # Assuming approx 40 paragraphs per page
        except Exception as e:
            print(f"Error reading {file_path}: {e}")
            return None

    def get_pptx_slide_count(file_path):
        try:
            prs = Presentation(file_path)
            return len(prs.slides)
        except Exception as e:
            print(f"Error reading {file_path}: {e}")
            return None

    # Initialize table headers
    headers = ["Sheet Name", "Folder Name", "Document Name", "Document Type", "Document Size (KB)", "No of Pages"]
    table = []

    # Check if the root folder exists
    if not os.path.exists(root_folder):
        print(f"Error: The folder '{root_folder}' does not exist.")
        return pd.DataFrame(table, columns=headers)

    for folder in os.listdir(root_folder):
        folder_path = os.path.join(root_folder, folder)
        
        # Only process directories
        if os.path.isdir(folder_path):
            for dirpath, _, filenames in os.walk(folder_path):
                sheet_name = os.path.basename(root_folder)  # Extract the sheet name based on the root folder

                for file in filenames:
                    file_path = os.path.join(dirpath, file)
                    file_extension = get_file_extension(file)
                    file_size = get_file_size(file_path)

                    # Handle page counts based on file type
                    if file_extension == '.pdf':
                        no_of_pages = get_pdf_page_count(file_path)
                    elif file_extension == '.docx':
                        no_of_pages = get_docx_page_count(file_path)
                    elif file_extension == '.pptx':
                        no_of_pages = get_pptx_slide_count(file_path)
                    else:
                        no_of_pages = 0  # For non-PDF, non-DOCX, and non-PPTX files, set page count to 0

                    # Add a row to the table
                    table.append([
                        sheet_name,  # Sheet Name
                        folder,
                        file,
                        file_extension,  # Use file extension as the document type
                        file_size,
                        no_of_pages
                    ])

    return pd.DataFrame(table, columns=headers)

# Set the main root folder path containing multiple root folders
main_root_folder = r"C:\Users\u1176867\OneDrive_2024-09-06 (2)"

# Initialize an Excel writer
output_path = r"C:\Users\u1176867\output.xlsx"
with pd.ExcelWriter(output_path) as writer:
    df_list = []
    for root_folder in os.listdir(main_root_folder):
        root_folder_path = os.path.join(main_root_folder, root_folder)
        
        # Only process directories
        if os.path.isdir(root_folder_path):
            df = list_all_files_summary(root_folder_path)
            df_list.append(df)
    
    # Concatenate all dataframes and write to a single sheet
    final_df = pd.concat(df_list, ignore_index=True)
    final_df.to_excel(writer, sheet_name="Summary", index=False)

print(f"The output has been saved to {output_path} with all data in a single sheet.")
